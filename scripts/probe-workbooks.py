"""
Reconnaissance probe for the two estimating workbooks.

Reports:
  - Sheet names + dimensions for each workbook.
  - Inv sheet row count + first/last item samples.
  - WC sheet row count + sample.
  - RateData sheet contents + B84 (rounding factor) value.
  - For the CNB workbook: every visible sheet's "Total" cell (heuristic: search
    for "Total" labels and grab the adjacent value), so we can sum to get the
    project acceptance target.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

REF = Path("reference/estimating")
BLANK = REF / "Sign365 LN Estimate Template - Blank.xlsx"
CNB = REF / "Sign365 LN Estimate - CNB.xlsx"


def header(title: str) -> None:
    print()
    print("=" * 72)
    print(title)
    print("=" * 72)


def summarise_workbook(path: Path, label: str) -> None:
    header(f"{label}: {path.name}")
    if not path.exists():
        print(f"  MISSING: {path}")
        return
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    print(f"  sheets ({len(wb_f.sheetnames)}):")
    for name in wb_f.sheetnames:
        ws = wb_f[name]
        state = ws.sheet_state
        print(f"    - {name!r:40s} state={state:8s} dims={ws.dimensions} "
              f"max_row={ws.max_row} max_col={ws.max_column}")
    return wb_f, wb_v


def show_inv(wb_f, wb_v) -> None:
    if "Inv" not in wb_f.sheetnames:
        print("  (no Inv sheet)")
        return
    header("Inv sheet")
    ws = wb_v["Inv"]
    # Show header row + a few samples + last row
    headers = [c.value for c in ws[1]]
    print(f"  headers: {headers}")
    print(f"  rows: {ws.max_row}")
    print("  first 3 data rows:")
    for r in range(2, 5):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        print(f"    r{r}: {row}")
    print("  last 2 data rows:")
    for r in range(ws.max_row - 1, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        print(f"    r{r}: {row}")


def show_wc(wb_f, wb_v) -> None:
    if "WC" not in wb_f.sheetnames:
        print("  (no WC sheet)")
        return
    header("WC sheet")
    ws = wb_v["WC"]
    print(f"  rows: {ws.max_row}, cols: {ws.max_column}")
    for r in range(1, min(ws.max_row, 15) + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        print(f"    r{r}: {row}")
    if ws.max_row > 15:
        print(f"    ... ({ws.max_row - 15} more rows)")


def show_ratedata(wb_f, wb_v) -> None:
    if "RateData" not in wb_f.sheetnames:
        print("  (no RateData sheet)")
        return
    header("RateData sheet")
    ws_v = wb_v["RateData"]
    ws_f = wb_f["RateData"]
    print(f"  rows: {ws_v.max_row}, cols: {ws_v.max_column}")
    # Show everything (RateData is small)
    for r in range(1, ws_v.max_row + 1):
        row_v = [ws_v.cell(r, c).value for c in range(1, ws_v.max_column + 1)]
        row_f = [ws_f.cell(r, c).value for c in range(1, ws_f.max_column + 1)]
        # Only print formulas if different from values
        if any(isinstance(v, str) and v.startswith("=") for v in row_f):
            print(f"    r{r} val: {row_v}")
            print(f"    r{r} fml: {row_f}")
        else:
            print(f"    r{r}: {row_v}")
    print()
    print(f"  B84 value: {ws_v['B84'].value!r}  formula: {ws_f['B84'].value!r}")


def find_totals(wb_f, wb_v, label: str) -> None:
    header(f"Total cells across all visible sheets — {label}")
    grand = 0.0
    for name in wb_f.sheetnames:
        ws = wb_f[name]
        if ws.sheet_state != "visible":
            continue
        ws_v = wb_v[name]
        # Search for cells containing "Total" (case-insensitive) and grab the
        # numeric value in the same row or adjacent cells.
        hits = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip().lower() == "total":
                    # Check the same row, columns to the right, for a number
                    for col in range(cell.column + 1, min(cell.column + 6, ws.max_column + 1)):
                        v = ws_v.cell(cell.row, col).value
                        if isinstance(v, (int, float)) and v != 0:
                            hits.append((f"{get_column_letter(cell.column)}{cell.row}",
                                         f"{get_column_letter(col)}{cell.row}", v))
                            break
        if hits:
            print(f"  {name!r}:")
            for label_cell, val_cell, v in hits:
                print(f"    {label_cell} 'Total' -> {val_cell} = {v}")
                if "project" not in name.lower():
                    pass  # do not auto-sum; just report
    # Also search for any cell labeled "Estimate Total" or "Project Total" or "Grand Total"
    print()
    print("  Special total labels (Grand/Project/Estimate Total):")
    for name in wb_f.sheetnames:
        ws = wb_f[name]
        if ws.sheet_state != "visible":
            continue
        ws_v = wb_v[name]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    s = cell.value.strip().lower()
                    if any(k in s for k in ("project total", "grand total", "estimate total", "job total")):
                        for col in range(cell.column + 1, min(cell.column + 6, ws.max_column + 1)):
                            v = ws_v.cell(cell.row, col).value
                            if isinstance(v, (int, float)):
                                print(f"    [{name}] {get_column_letter(cell.column)}{cell.row} "
                                      f"{cell.value!r} -> {get_column_letter(col)}{cell.row} = {v}")
                                break


def main() -> None:
    blank_res = summarise_workbook(BLANK, "BLANK template")
    cnb_res = summarise_workbook(CNB, "CNB filled job")
    if blank_res:
        wb_f, wb_v = blank_res
        show_inv(wb_f, wb_v)
        show_wc(wb_f, wb_v)
        show_ratedata(wb_f, wb_v)
    if cnb_res:
        wb_f, wb_v = cnb_res
        find_totals(wb_f, wb_v, "CNB")


if __name__ == "__main__":
    main()
