"""Dump CNB BCI + BCL sheets to compute the actual project total."""
from pathlib import Path
from openpyxl import load_workbook

CNB = Path("reference/estimating") / "Sign365 LN Estimate - CNB.xlsx"

wb_v = load_workbook(CNB, data_only=True)
wb_f = load_workbook(CNB, data_only=False)

print("=" * 72)
print("BCI sheet (materials export)")
print("=" * 72)
ws = wb_v["BCI"]
ws_f = wb_f["BCI"]
print(f"  headers: {[c.value for c in ws[1]]}")
mat_total = 0.0
mat_lines = 0
for r in range(2, ws.max_row + 1):
    row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if any(v is not None and v != 0 and v != "" for v in row):
        # Job#, Task#, Item#, Desc, Unit Cost, Profit%, Unit Price, Qty
        job_no, task_no, item_no, desc, unit_cost, profit_pct, unit_price, qty = (row + [None]*8)[:8]
        if isinstance(unit_price, (int, float)) and isinstance(qty, (int, float)) and qty > 0:
            line_total = unit_price * qty
            mat_total += line_total
            mat_lines += 1
            print(f"  r{r}: Job={job_no} Task={task_no} Item={item_no} "
                  f"Desc={desc!r} UC={unit_cost} P%={profit_pct} "
                  f"UP={unit_price} Qty={qty} -> ${line_total:.4f}")
print(f"\n  BCI material lines: {mat_lines}")
print(f"  BCI material subtotal: ${mat_total:.4f}")

print()
print("=" * 72)
print("BCL sheet (labor export)")
print("=" * 72)
ws = wb_v["BCL"]
print(f"  headers: {[c.value for c in ws[1]]}")
# Need work-code -> rate map
wc = wb_v["WC"]
rate_by_code = {}
for r in range(6, wc.max_row + 1):
    code = wc.cell(r, 1).value
    rate = wc.cell(r, 3).value
    if isinstance(code, (int, float)) and isinstance(rate, (int, float)):
        rate_by_code[int(code)] = float(rate)
print(f"  rate lookup (codes -> hourly rate): {rate_by_code}")
lab_total = 0.0
lab_lines = 0
for r in range(2, ws.max_row + 1):
    row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if any(v is not None and v != 0 and v != "" for v in row):
        job_no, resource_no, run_time = (row + [None]*3)[:3]
        if isinstance(run_time, (int, float)) and run_time > 0:
            rate = rate_by_code.get(int(resource_no) if isinstance(resource_no, (int, float)) else -1, 0)
            line_total = run_time * rate
            lab_total += line_total
            lab_lines += 1
            print(f"  r{r}: Job={job_no} Resource={resource_no} hrs={run_time} "
                  f"@${rate}/hr -> ${line_total:.4f}")
print(f"\n  BCL labor lines: {lab_lines}")
print(f"  BCL labor subtotal: ${lab_total:.4f}")

print()
print("=" * 72)
print(f"CNB PROJECT TOTAL = ${mat_total + lab_total:.4f}")
print("=" * 72)
