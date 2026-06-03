"""Inspect a single piece sheet's full layout so we can identify the input
block, the material block, the labor block, and the piece total cell.

Usage: python scripts/probe-piece-layout.py "Sheet Name"
"""
from pathlib import Path
import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

REF = Path("reference/estimating")
BLANK = REF / "Sign365 LN Estimate Template - Blank.xlsx"
CNB = REF / "Sign365 LN Estimate - CNB.xlsx"


def dump(workbook_path: Path, sheet_name: str, max_col: int = 12, max_row: int = 70) -> None:
    print(f"\n=== {workbook_path.name} / {sheet_name!r} ===")
    wb_v = load_workbook(workbook_path, data_only=True)
    wb_f = load_workbook(workbook_path, data_only=False)
    if sheet_name not in wb_v.sheetnames:
        print(f"  (sheet not found)")
        return
    ws_v = wb_v[sheet_name]
    ws_f = wb_f[sheet_name]
    for r in range(1, min(ws_v.max_row, max_row) + 1):
        cells = []
        for c in range(1, min(ws_v.max_column, max_col) + 1):
            v = ws_v.cell(r, c).value
            f = ws_f.cell(r, c).value
            if v is None and f is None:
                cells.append("")
                continue
            label = get_column_letter(c)
            if isinstance(f, str) and f.startswith("="):
                # Show formula and computed value.
                cells.append(f"{label}={v!r:.30s} <{f[:55]}>")
            else:
                cells.append(f"{label}={v!r}")
        if any(c for c in cells):
            print(f"  r{r:>2}: " + " | ".join(c for c in cells if c))


def main() -> None:
    sheet = sys.argv[1] if len(sys.argv) > 1 else "Routed panel shapes"
    dump(CNB, sheet)
    print()
    print("(blank for comparison:)")
    dump(BLANK, sheet)


if __name__ == "__main__":
    main()
