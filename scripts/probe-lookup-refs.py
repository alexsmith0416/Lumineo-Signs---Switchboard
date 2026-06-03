"""Find which RateData! cells the piece sheets reference via VLOOKUP, so we
can confirm the correct column for the increment-table rates (new vs old)."""
from pathlib import Path
from openpyxl import load_workbook
import re

BLANK = Path("reference/estimating") / "Sign365 LN Estimate Template - Blank.xlsx"
wb_f = load_workbook(BLANK, data_only=False)

# Look for VLOOKUP formulas that point at RateData and extract the range.
pat = re.compile(r"VLOOKUP\([^,]+,\s*RateData!([A-Z]+\$?\d+:[A-Z]+\$?\d+)", re.IGNORECASE)
hits = {}
for name in wb_f.sheetnames:
    if wb_f[name].sheet_state != "visible":
        continue
    ws = wb_f[name]
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "RateData" in cell.value:
                for m in pat.finditer(cell.value):
                    rng = m.group(1)
                    hits.setdefault(rng, []).append((name, cell.coordinate, cell.value[:120]))

for rng, occ in sorted(hits.items()):
    print(f"\nRateData!{rng}  ({len(occ)} refs)")
    for sheet, coord, val in occ[:3]:
        print(f"  [{sheet}] {coord}: {val}")
